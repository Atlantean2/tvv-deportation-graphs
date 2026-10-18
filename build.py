"""Builds the TechVestVentures Deportation Graphs site from focus/immigration (pptx + docx + CSVs)."""
import os,re,sys,shutil,glob
sys.path.insert(0,'/workspace/risk-watcher/graphs-site');os.chdir(os.path.dirname(os.path.abspath(__file__)))
from common import *
from pptx import Presentation
T="TechVestVentures Deportation Graphs";F="/workspace/risk-watcher/focus/immigration"
P=Presentation(glob.glob(F+"/TVV Focus Immigration Graphs Collection *.pptx")[0])
shutil.rmtree("issues",ignore_errors=True);os.makedirs("issues/img");os.makedirs("issues/files")
secs={};heads={};summ={};allsrc=[];order=[]
for i,s in enumerate(P.slides,1):
    tx=[sh.text_frame.text.strip() for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
    ttl=s.shapes.title.text.strip() if s.shapes.title else (tx[0].split("\n")[0] if tx else "")
    n=s.notes_slide.notes_text_frame.text if s.has_notes_slide else ""
    m=re.match(r"(Q\d)\. (.*)",ttl)
    if m:
        heads[m[1]]=m[2];summ[m[1]]=[t for t in tx if t!=ttl];continue
    for u in re.findall(r"https://[^\s)\]>,]+",n):
        if u not in allsrc: allsrc.append(u)
    m=re.match(r"(Q\d): (.*)",ttl)
    if not m: continue
    pics=[sh for sh in s.shapes if sh.shape_type==13];pic=max(pics,key=lambda x:x.width*x.height)
    f=dict(l.split(": ",1) for l in n.splitlines() if ": " in l and not l.startswith("Source"))
    urls=re.findall(r"^Source: (https://\S+)",n,re.M)
    title=f.get("Title",m[2]);key="TVV" if title.startswith("TVV derived") else m[1]
    cap=f"{f.get('Publisher','')}. {f.get('What it counts','')}".strip(". ")
    lst=secs.setdefault(key,[]);fn=f"{key}-{len(lst)+1:02d}.webp";save_img(pic.image.blob,"issues/img/"+fn)
    lst.append(figure("img/"+fn,f"{m[1]}: {title}" if key=="TVV" else title,cap,urls[:2]))
tabs=[]
def w(k,lab,body): open(f"issues/{k}.html","w").write(page(T,k,body));tabs.append((k,lab))
# topical material
for f in glob.glob(F+"/*.docx")+glob.glob(F+"/data/*.csv"): shutil.copy(f,"issues/files/")
L=lambda p:f'<a href="files/{E(os.path.basename(p))}" download>{E(os.path.basename(p))}</a>'
b=['<h2 class="desk">Topical material</h2><p>Research notes (Word document): '+" ".join(L(p) for p in glob.glob(F+"/*.docx"))+'</p>',
   '<h2 class="desk">Section questions (from the Graphs Collection)</h2><ul class="items">']
for q in sorted(heads):
    b.append(f'<li><a href="#{q}">{q}. {E(heads[q])}</a><ul>'+"".join(f"<li>{E(x)}</li>" for t in summ[q] for x in t.split("\n") if x.strip())+"</ul></li>")
b.append('</ul><h2 class="desk">Data files (CSV)</h2><ul class="items">'+"".join(f"<li>{L(p)}</li>" for p in sorted(glob.glob(F+"/data/*.csv")))+"</ul>")
w("material","Topical material","".join(b))
for q in sorted(k for k in secs if k!="TVV"):
    w(q,f"{q}{' Self-deportation' if q=='Q6' else ''}",f'<h2 class="desk">{q}. {E(heads.get(q,""))}</h2>'+"\n".join(secs[q]))
w("TVV","TVV derived charts",'<h2 class="desk">TVV derived charts (built from primary government data)</h2>'+"\n".join(secs["TVV"]))
w("sources","Sources",f'<h2 class="desk">Sources ({len(allsrc)} links)</h2><ul class="src">'+"".join(f'<li><a href="{E(u)}">{E(u)}</a></li>' for u in allsrc)+"</ul>")
tabs=[t for t in tabs if t[0]!="material"]+[t for t in tabs if t[0]=="material"];tabs=tabs[:-2]+[tabs[-1],tabs[-2]]
index(T,tabs,"Sections");print({k:len(v) for k,v in secs.items()},len(allsrc))
